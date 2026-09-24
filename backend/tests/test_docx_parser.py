"""Parser DOCX: ordem do corpo e tachado."""

from pathlib import Path

from docx import Document as DocxDocument

from app.services.parser.docx_parser import parse_docx
from app.services.parser.legal_marks import normalize_legal_text


def test_docx_preserva_ordem_e_tachado(tmp_path: Path):
    doc = DocxDocument()
    doc.add_paragraph("Antes da tabela")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "celula"
    para = doc.add_paragraph()
    para.add_run("vigente ")
    struck = para.add_run("antigo")
    struck.font.strike = True
    para.add_run(" depois")

    path = tmp_path / "ordem.docx"
    doc.save(path)

    text, pages = parse_docx(path)
    assert text.index("Antes da tabela") < text.index("celula")
    assert text.index("celula") < text.index("vigente")
    assert "[TABELA]" in text
    assert "[TACHADO]antigo[/TACHADO]" in text
    assert pages[0]["page"] == 1

    normalized = normalize_legal_text(text)
    assert "antigo" not in normalized.vigente
    assert any(mark.kind == "strikethrough" for mark in normalized.marks)
